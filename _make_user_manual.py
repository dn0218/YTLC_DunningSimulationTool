# -*- coding: utf-8 -*-
"""Generate 'Dunning Tool User Manual Guide.docx' (English user manual + SOP).

Chapter structure mirrors: preprod tools/docs/USER_MANUAL_en.md
  1. What This Tool Is
  2. Quick Start
  3. Step-by-Step SOPs
  4. Troubleshooting Mini-FAQ
  5. Issue Feedback Process
  6. Safety Rules
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "Dunning Tool User Manual Guide.docx"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue
GRAY_FILL = "F2F2F2"
NOTE_FILL = "FFF2CC"                   # light amber
WARN_FILL = "FCE4E4"                   # light red


# ---------------------------------------------------------------- helpers
def shade(paragraph, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def para(doc, text="", bold_prefix=None, size=10.5, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(size)
    if text:
        r = p.add_run(text)
        r.italic = italic
        r.font.size = Pt(size)
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def code_block(doc, text, fill=GRAY_FILL):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9)
        shade(p, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def note_box(doc, label, text, fill=NOTE_FILL):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(label + "  ")
    r.bold = True
    p.add_run(text)
    shade(p, fill)


def table(doc, headers, rows, widths=None, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        shade(hdr[i].paragraphs[0], "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def step(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step %d. " % n)
    r.bold = True
    r.font.color.rgb = ACCENT
    p.add_run(text)


# ---------------------------------------------------------------- document
doc = Document()
doc.core_properties.title = "Dunning Tool User Manual Guide"
doc.core_properties.subject = "Dunning Simulation Tool v4 - User Manual & SOP (Test Environment)"
doc.core_properties.author = "Dunning Tool Maintainers"

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
for i in (1, 2, 3):
    hs = doc.styles["Heading %d" % i]
    hs.font.color.rgb = ACCENT

# footer
fp = sec.footer.paragraphs[0]
fp.text = "Dunning Simulation Tool v4 — User Manual & SOP  |  Test Environment Only"
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in fp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

# ---------------- cover block
tp = doc.add_heading("Dunning Simulation Tool", 0)
tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub = doc.add_paragraph()
r = sub.add_run("User Manual Guide & Standard Operating Procedures (SOP)")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = ACCENT

table(doc,
      ["Item", "Value"],
      [["Tool name", "Dunning Simulation Tool v4 (DunningSimulationTool.exe / python main.py)"],
       ["Applies to", "Telecom billing TEST environment (Collections / Dunning module) — NEVER production"],
       ["Document version", "v1.0"],
       ["Related internal docs", "交接文档/00–06 (Chinese, for maintainers): business SQL logic, config, ops, deployment, maintenance"],
       ],
      widths=[4.2, 12.6])

doc.add_heading("Contents", level=1)
for line in ["1. What This Tool Is",
             "2. Quick Start",
             "3. Step-by-Step SOPs  (SOP-1 … SOP-4 + level reference tables)",
             "4. Troubleshooting Mini-FAQ",
             "5. Issue Feedback Process",
             "6. Safety Rules"]:
    p = doc.add_paragraph(line)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.6)

# ================= 1. What This Tool Is
doc.add_heading("1. What This Tool Is", level=1)
para(doc,
     "The Dunning Simulation Tool is a desktop application for the billing TEST environment. "
     "In production, a customer who does not pay walks through the dunning (collections) levels one "
     "by one — reminder, barring, suspension, IMEI/CTOS blacklist, termination — as real calendar "
     "days pass. Waiting 105 real days to test one action is impossible. This tool makes the "
     "account LOOK N days overdue by adjusting its billing-cycle data, then fires the very same "
     "backend jobs production uses (DebitActionJob / DebitActionDealJob). The whole chain runs in "
     "minutes, from one window.")

para(doc,
     "You never log in to the database, SSH, or the job server yourself. You provide an account, "
     "click the buttons in order, and read the result in the log panel. The workflow is always the "
     "same three steps:")

table(doc,
      ["Step", "You input", "Button", "What it does", "Success looks like"],
      [["1. Get Account",
        "Account ID or account number",
        "Get Account",
        "Finds the account, detects its customer type (MASSIVE / ENTERPRISE / HYBRID) and fills the level dropdown",
        "Blue label USER_ACCT: <id> (<type>); status bar shows Done"],
       ["2. Initial Setup",
        "— (uses the account from step 1)",
        "Run Initial Setup",
        "Puts the account into a valid starting state (debt check, auto-pay off, billing cycle prepared, jobserver restarted) and generates the first (Level 0) overdue instance",
        "Dialog \u201cInitial setup completed.\u201d"],
       ["3. Run Simulation",
        "Level selection",
        "Run Simulation",
        "Ages the account to the selected level and lets the real jobs execute that level\u2019s actions (SMS/e-mail, barring, suspension, IMEI/CTOS, termination)",
        "Dialog \u201cLevel \u2026 simulation succeeded!\u201d and EXEC_TIME in the log"],
       ],
      widths=[2.6, 3.0, 2.6, 5.6, 4.0])

para(doc, "Three universal rules that apply to every step:", bold_prefix="")
bullet(doc, "before starting another action — disabled buttons mean the tool is working (job polling can take up to 60 seconds).",
       bold_prefix="Wait for buttons to re-enable ")
bullet(doc, "Initial Setup writes directly to the test database (attribute fix, auto-pay disabled, billing-cycle dates shifted, scan-scope config changed). Never point the tool at production.",
       bold_prefix="Use test accounts only. ")
bullet(doc, "until the operation is confirmed successful — it is the evidence required for troubleshooting (Section 5).",
       bold_prefix="Keep the log panel content ")

# ================= 2. Quick Start
doc.add_heading("2. Quick Start", level=1)

doc.add_heading("2.1 Start the tool", level=2)
bullet(doc, "double-click DunningSimulationTool.exe. No installation needed.", bold_prefix="Packaged build: ")
bullet(doc, "python main.py  (see internal doc 04 for environment setup).", bold_prefix="From source (maintainers only): ")

doc.add_heading("2.2 First launch — one-time setup", level=2)
step(doc, 1, "The tool reads config.json from the folder next to the exe. If it is missing, the tool writes a default template and shows a config error.")
step(doc, 2, "Normally you do NOT touch this file — it ships pre-configured for the test environment.")
step(doc, 3, "Only if your maintainer asks (e.g. the environment moved): edit it in Notepad (DB / SSH / API sections), save, restart the tool.")
step(doc, 4, "config.json contains PLAINTEXT credentials — never share, screenshot, or commit it (see Section 6).")
code_block(doc, '''config.json sections:
  "db"      -> host / port / service_name / user / password      (Oracle, test CC schema)
  "ssh"     -> host / port / username / password / timeouts      (jobserver host)
  "api"     -> base_url / timeout / retries                      (job trigger HTTP)
  "general" -> poll / wait timings in seconds (defaults are fine)''')

doc.add_heading("2.3 What you need", level=2)
bullet(doc, "A PC that can reach the test network (Oracle 10.26.128.32:11521, SSH/API host 10.26.128.43).")
bullet(doc, "A TEST account identifier: Account ID or account number. The account should carry unpaid charges (debt) — that is the whole point.")
bullet(doc, "For result verification: access to Billing Center Portal (Dunning Instance Management, IMEI Blacklist Management, Account Special Management) — optional but recommended.")

doc.add_heading("2.4 The window at a glance", level=2)
table(doc,
      ["Area", "What is there"],
      [["Top row", "Account input + Get Account button, USER_ACCT result label, Language combo (English / Chinese)"],
       ["Second row", "Run Initial Setup button"],
       ["Third row", "Customer Type combo (MASSIVE / ENTERPRISE / HYBRID), Level combo, Run Simulation button"],
       ["Status line", "Status: Ready / Running\u2026 / Done / Error, and a Clear Log button"],
       ["Log panel", "Every DB write, remote command, API call and check is printed here with timestamps — keep it until you are done"]],
      widths=[3.2, 13.6])

# ================= 3. SOPs
doc.add_heading("3. Step-by-Step SOPs", level=1)
para(doc, "Run the SOPs strictly in order: SOP-1 once per account, SOP-2 once per account, "
          "then SOP-3 repeatedly — once per level you want to test. SOP-4 closes the session.")

doc.add_heading("SOP-1 — Look Up an Account (Get Account)", level=2)
para(doc, "Purpose: resolve the account and detect its customer type & available levels.", italic=True)
step(doc, 1, "Type the Account ID or the account number into the \u201cAccount ID or Number\u201d box.")
step(doc, 2, "Press Enter or click Get Account.")
step(doc, 3, "Read the results: the blue label shows USER_ACCT: <id> (<type>); the Customer Type combo auto-selects MASSIVE / ENTERPRISE / HYBRID; the Level dropdown fills with that group\u2019s levels.")
para(doc, "Expected result: status bar shows Done. Log shows \u201cFound USER_ACCT \u2026\u201d and \u201cRule group resolved to \u2026\u201d.",
     italic=True)
note_box(doc, "Note:", "The customer type is detected automatically, but you can switch it manually "
         "with the combo (e.g. to preview another group\u2019s levels). Initial Setup uses the type shown "
         "at the moment you click it.", NOTE_FILL)

doc.add_heading("SOP-2 — Run Initial Setup (prepare the account)", level=2)
para(doc, "Purpose: put the account into a valid testable state and generate the Level-0 overdue "
          "instance using the real daily job.", italic=True)
note_box(doc, "CAUTION:", "This SOP writes to the test database: fixes the customer attribute "
         "(ATTR_ID 907207), disables auto-pay signing, sets the billing cycle to closed with a "
         "DEBT_DATE of SYSDATE-3, creates TMP_DUNNING, repoints the job scan scope, and restarts "
         "the jobserver. Test accounts only!", WARN_FILL)
step(doc, 1, "Complete SOP-1 first (an account must be shown in the blue label).")
step(doc, 2, "Click Run Initial Setup. Buttons become disabled — this is normal.")
step(doc, 3, "Watch the log; the tool runs: attribute check \u2192 debt check \u2192 billing cycle state check (retries if \u2018B\u2019 = billing in progress) \u2192 auto-pay off \u2192 TMP_DUNNING prepared \u2192 jobserver restart over SSH \u2192 DebitActionJob triggered \u2192 waiting for the overdue instance.")
step(doc, 4, "Success: dialog \u201cInitial setup completed.\u201d; the log shows the generated INST_ID and the matched Level 0 rule/action ID.")
para(doc, "Duration: typically 1\u20133 minutes (instance polling every 3 s up to 60 s; billing-state "
          "retries every 10 s up to 3 times).", italic=True)
note_box(doc, "If it fails:", "See Section 4, fix the cause, and simply run SOP-2 again — it is safe "
         "to rerun and will rebuild the state.", WARN_FILL)

doc.add_heading("SOP-3 — Simulate a Dunning Level (Run Simulation)", level=2)
para(doc, "Purpose: make the account look N days overdue so the real jobs execute exactly that "
          "level\u2019s actions.", italic=True)
step(doc, 1, "Check the Customer Type combo (auto-selected in SOP-1; change only if you know why).")
step(doc, 2, "Pick the Level in the dropdown. Each entry shows: rule name [Rule ID, offset days].")
step(doc, 3, "Click Run Simulation and watch the log: instance reset/regeneration \u2192 DEBT_DATE / DUE_DATE update \u2192 prerequisite rule chain (when needed) \u2192 action execution \u2192 EXEC_TIME check.")
step(doc, 4, "Success: dialog \u201cLevel \u2026 simulation succeeded!\u201d, log lines \u201cEXEC_TIME generated: \u2026\u201d and \u201cDEBIT_ACTION_HIS.STATE = \u2026\u201d. Keep this log.")
step(doc, 5, "Verify the business effect if needed: Dunning Instance Management (instance Running / rule actions), IMEI Blacklist Management (Level 4+), Account Special Management (CTOS entries), or the SMS/e-mail received by the test MSISDN.")
step(doc, 6, "Repeat from step 2 for the next level. Picking a LOWER level later works too — the tool automatically resets the instance and walks the chain again.")
note_box(doc, "Normal behaviour:", "For Level 0 the tool temporarily shifts CYCLE_END_DATE and "
         "restores it after the run — a short \u201crestoring BILLING_CYCLE \u2026\u201d pause at the end is "
         "expected. A \u201cclamping to \u2026\u201d message means an unparsable offset was safely corrected.")

doc.add_heading("SOP-4 — After Testing (clean exit)", level=2)
step(doc, 1, "Note which account(s) you tested and the levels reached.")
step(doc, 2, "Close the window — the DB/SSH connections close automatically.")
step(doc, 3, "Inform your maintainer / environment owner that testing was done: the job scan scope "
             "(DEBIT_DEAL_SEARCH_SQL, CONFIG_ITEM_ID 46/234) is intentionally left pointing at TMP_DUNNING "
             "and is NOT restored by the tool. Daily jobs on this environment will only scan TMP_DUNNING "
             "until it is reverted.")
step(doc, 4, "If a fully clean environment state is required, ask the maintainer (manual restore steps are in internal doc 05, \u00a77).")

doc.add_heading("3.5 Level reference tables (what each level means)", level=2)
para(doc, "The dropdown labels come live from the database (DEBIT_RULE), so they always match the "
          "current configuration. The tables below are the normal timeline for orientation.")
table(doc,
      ["Level", "Timing", "Action (MASSIVE, group 4)"],
      [["0", "5 days before due date", "Reminder (SMS / e-mail)"],
       ["1", "Due date + 1 day", "Barring (outbound blocked)"],
       ["2", "+ 10 days", "Notification"],
       ["3", "+ 20 days", "Notification"],
       ["4", "+ 30 days", "Suspension & IMEI blacklist"],
       ["5", "+ 40 days", "Pre-CTOS notification"],
       ["6", "+ 45 days", "CTOS blacklist"],
       ["7", "+ 60 days", "Payment reminder"],
       ["8", "+ 95 days", "Final reminder"],
       ["9", "+ 105 days", "Termination"]],
      widths=[1.6, 4.6, 10.6])
table(doc,
      ["Level", "Timing", "Action (ENTERPRISE, group 5)"],
      [["0", "25 days before due date", "Reminder (SMS / e-mail)"],
       ["1", "Due date + 1 day", "Barring"],
       ["2", "+ 20 days", "Pre-suspension notification"],
       ["3", "+ 30 days", "Suspension"],
       ["4", "+ 80 days", "Pre-termination notification"],
       ["5", "+ 105 days", "Termination"]],
      widths=[1.6, 4.6, 10.6])
table(doc,
      ["Level", "Timing", "Action (HYBRID, group 14)"],
      [["0", "Due date + 31 days", "Notification"],
       ["1", "+ 53 days", "Final notification"],
       ["2", "+ 60 days", "Termination"]],
      widths=[1.6, 4.6, 10.6])

# ================= 4. FAQ
doc.add_heading("4. Troubleshooting Mini-FAQ", level=1)
para(doc, "Try these before reporting an issue — most reports stop here.")
table(doc,
      ["Symptom", "What to do"],
      [["Tool won\u2019t start; message about config.json / Config",
        "First-launch behaviour: fill in config.json next to the exe (Section 2.2) and restart. If it persists, ask your maintainer — do not guess credentials."],
       ["\u201cDatabase connection failed\u201d at startup",
        "You are probably off the test network (VPN). Retry after reconnecting. If colleagues are fine and you are not, it is your connection; if everyone fails, it is the environment — report as an environment issue, not a tool bug."],
       ["\u201cAccount not found.\u201d",
        "Check the identifier: try the other form (account number vs account ID). Not a bug."],
       ["\u201cAccount has no debt, please perform debit adjustment.\u201d",
        "Business validation, not a bug: the account must carry unpaid charges. Add debt to the test account or pick another one."],
       ["\u201cNo billing cycle found.\u201d",
        "The account has no closed bill items. Use a properly billed test account."],
       ["\u201cBilling Cycle still in \u2018B\u2019 state after 3 checks, aborting setup.\u201d",
        "Billing is running on the server right now. Wait a few minutes and rerun SOP-2."],
       ["\u201cDebitActionJob failed to generate overdue instance \u2026\u201d",
        "The job did not produce the instance within the wait window. Retry SOP-2 once; if it repeats, keep the log and report it (Section 5)."],
       ["\u201cEXEC_TIME field not updated, simulation did not complete successfully.\u201d",
        "The action was generated but failed to execute. The log usually shows the history error message (ERROR_MSG) — attach the log to your report."],
       ["\u201cFailed to start Jobserver after retries.\u201d",
        "SSH / jobserver problem on the server. Report as an environment issue with the log."],
       ["Buttons grey / window seems \u201cfrozen\u201d",
        "Normal: built-in waits (a few seconds between DB writes; job polling up to 60 s). Wait for the buttons to re-enable. Also check whether a dialog is hiding behind the window."],
       ["Level dropdown is empty",
        "No actionable levels were loaded for the rule group (database rule configuration). Report it with the startup log."],
       ["Nothing happened after I clicked twice quickly",
        "The tool blocks parallel operations on purpose. One operation at a time (Section 6)."]],
      widths=[6.2, 10.6])
para(doc, "Still stuck? \u2192 Go to Section 5 and follow the feedback process.")

# ================= 5. Feedback
doc.add_heading("5. Issue Feedback Process", level=1)
para(doc, "Read this before reporting anything. Reports that do not follow this process will NOT "
          "be accepted. A complete report is answered in one batch cycle; an incomplete one costs "
          "everyone a week of back-and-forth.")

doc.add_heading("5.1 Where to submit", level=2)
bullet(doc, "Submit via the Teams feedback channel to the designated tester of your team. Do NOT send reports directly to developers\u2019 private chats.")
bullet(doc, "Testers collect all reports during the week and submit them as one batch. Maintainers process batches 1\u20132 times per week.")
bullet(doc, "Consequence: a report filed on Monday may be answered at the weekly batch — plan your testing accordingly.")

doc.add_heading("5.2 Issue template (copy, fill in, attach — missing fields = not accepted)", level=2)
code_block(doc, '''## Issue Report (all fields required)
- [ ] Tool version - exe "Date modified" (right-click
      DunningSimulationTool.exe -> Properties -> Details):
- [ ] Section:  [ ] Startup/Config  [ ] Get Account
                [ ] Initial Setup   [ ] Run Simulation
- [ ] Account ID used:
- [ ] Customer type & level (if Simulation):
- [ ] Steps to reproduce (1 / 2 / 3 ..., exact inputs):
- [ ] Expected result:
- [ ] Actual result (exact error text / dialog title):
- [ ] Console log screenshot (REQUIRED - the log panel,
      scrolled to the error, not just the input screen):
- [ ] Did you check Section 4 FAQ first?  [ ] Yes
- [ ] Reporter + date:''')
para(doc, "How to fill \u201cTool version\u201d: right-click DunningSimulationTool.exe \u2192 Properties \u2192 "
          "Details \u2192 Date modified — type that date. Two bugs with different dates are different "
          "builds; without it we cannot reproduce.", italic=True)

doc.add_heading("5.3 Acceptance criteria (a report is accepted only when ALL hold)", level=2)
step(doc, 1, "Submitted through the Teams channel to the designated tester (not private chats).")
step(doc, 2, "Every template field filled in — including the version date and the FAQ checkbox.")
step(doc, 3, "Console log screenshot attached, showing the error (not just the input screen).")
step(doc, 4, "Reproducible: the steps, with the exact account ID used, lead to the problem.")
step(doc, 5, "Not already reported — check the channel\u2019s pinned \u201cKnown Issues\u201d list first.")

doc.add_heading("5.4 NOT accepted (any one of these = rejected, no reply)", level=2)
step(doc, 1, "Verbal / phone / private-chat reports — no record, no tracking.")
step(doc, 2, "Missing console log or screenshots (\u201cit just failed\u201d is not a report).")
step(doc, 3, "Missing steps to reproduce, or steps that cannot be repeated.")
step(doc, 4, "\u201cIt doesn\u2019t work\u201d with no error text or expected-vs-actual comparison.")
step(doc, 5, "Issue already listed in the pinned Known Issues list.")
step(doc, 6, "Environment problems (DB down, network off, VPN) reported as tool bugs — see Section 4.")

doc.add_heading("5.5 Status flow & response times", level=2)
table(doc,
      ["Status", "Meaning", "Who acts", "Expected within"],
      [["New", "Batch received by maintainer", "Maintainer", "Next batch day (1\u20132\u00d7/week)"],
       ["Triaged", "Confirmed valid, priority set", "Maintainer", "Same batch cycle, with a first reply"],
       ["Info needed", "Report incomplete \u2192 back to reporter", "You", "Reply within 3 working days, or the issue is closed"],
       ["Fixing", "Fix in progress", "Maintainer", "Announced ETA in the channel"],
       ["Fixed", "New build posted in the channel", "Maintainer", "\u2014"],
       ["Closed", "Verified by reporter or auto-closed", "You", "Verify in the new build within 5 working days"]],
      widths=[2.6, 5.6, 3.0, 5.6])
para(doc, "One report = one issue. If you found three problems, submit three copies of the template.")

# ================= 6. Safety
doc.add_heading("6. Safety Rules", level=1)
step(doc, 1, "This tool writes directly to the database it is configured for. Never point config.json at production; never run it with production credentials.")
step(doc, 2, "config.json contains PLAINTEXT credentials — do not share, upload, or screenshot it.")
step(doc, 3, "One operation at a time. Wait for buttons to re-enable; never close the tool in the middle of Initial Setup or Simulation.")
step(doc, 4, "Keep the log until the result is confirmed — and attach it to every issue report (Section 5).")
step(doc, 5, "Remember the side effects: auto-pay is disabled, the billing cycle is shifted, and the job scan scope is left at TMP_DUNNING after testing (SOP-4). Coordinate with everyone who shares the test environment.")
step(doc, 6, "When unsure, ask before clicking — in the Teams channel, before the operation, not after.")

doc.save(OUT)
print("saved:", OUT)
